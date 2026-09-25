from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_text(cell, text: Any, *, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run("" if text is None else str(text))
    run.bold = bold
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = __import__("docx").shared.RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _add_section_heading(document: Document, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = __import__("docx").shared.RGBColor.from_string(NAVY)


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def _set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _format_table(table) -> None:
    table.style = "Table Grid"
    table.autofit = True
    for row in table.rows:
        for cell in row.cells:
            _set_cell_margins(cell)


def _add_header_footer(document: Document, project_name: str, document_type: str) -> None:
    for section in document.sections:
        header = section.header
        hp = header.paragraphs[0]
        hp.text = ""
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = hp.add_run(f"{project_name}  |  {document_type}")
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = __import__("docx").shared.RGBColor.from_string("6B7280")

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run(
            "AV Intelligence Assistant  |  Field verification required for unconfirmed items"
        )
        run.font.name = "Aptos"
        run.font.size = Pt(7.5)
        run.font.color.rgb = __import__("docx").shared.RGBColor.from_string("7F7F7F")


def _add_document_control_table(
    document: Document,
    *,
    project_name: str,
    client: str,
    location: str,
    rooms: list[str],
    status: str,
    document_type: str,
) -> None:
    table = document.add_table(rows=0, cols=2)
    _format_table(table)
    rows = [
        ("Document", document_type),
        ("Project", project_name),
        ("Client", client),
        ("Location", location),
        ("Rooms / Areas", ", ".join(rooms)),
        ("Status", status),
    ]
    for label, value in rows:
        if not value:
            continue
        cells = table.add_row().cells
        _set_cell_text(cells[0], label, bold=True, color=WHITE)
        _shade_cell(cells[0], NAVY)
        _set_cell_text(cells[1], value)


def _normalize_required_photos(checklist_data: dict[str, Any]) -> list[str]:
    normalized = []
    for item in checklist_data.get("required_photos", []) or []:
        if isinstance(item, dict):
            subject = (
                item.get("photo_subject")
                or item.get("subject")
                or item.get("item")
                or ""
            )
            location = item.get("location_direction") or ""
            related = item.get("related_item") or ""
            text = " | ".join(
                str(value).strip()
                for value in (subject, location, related)
                if str(value or "").strip()
            )
        else:
            text = str(item or "").strip()
        if text:
            normalized.append(text)
    return normalized


def _normalized_checklist_sections(checklist_data: dict[str, Any]) -> list[dict[str, Any]]:
    sections = []

    for section in _normalized_checklist_sections(checklist_data):
        if not isinstance(section, dict):
            continue
        name = str(section.get("section") or "").strip()
        items = []
        for item in section.get("items", []) or []:
            if not isinstance(item, dict):
                continue
            items.append(
                {
                    "item": item.get("item") or item.get("inspection_item") or "",
                    "status": item.get("status") or "VERIFY",
                    "notes": item.get("notes") or item.get("notes_photo") or "",
                    "photo_ref": item.get("photo_ref") or "",
                }
            )
        if name and items:
            sections.append({"section": name, "items": items})

    if sections:
        return sections

    for section in checklist_data.get("inspection_sections", []) or []:
        if not isinstance(section, dict):
            continue
        name = str(
            section.get("section_title")
            or section.get("section")
            or "Inspection"
        ).strip()
        items = []
        for item in section.get("items", []) or []:
            if not isinstance(item, dict):
                continue
            items.append(
                {
                    "item": item.get("inspection_item") or item.get("item") or "",
                    "status": item.get("status") or "VERIFY",
                    "notes": item.get("notes_photo") or item.get("notes") or "",
                    "photo_ref": item.get("photo_ref") or "",
                }
            )
        if name and items:
            sections.append({"section": name, "items": items})

    return sections


def _visual_inspection_summary(checklist_data: dict[str, Any]) -> dict[str, Any]:
    visual = checklist_data.get("visual_inspection")
    return visual if isinstance(visual, dict) else {}


def _status_summary(checklist_data: dict[str, Any]) -> dict[str, int]:
    counts = {"PASS": 0, "VERIFY": 0, "ACTION": 0, "N/A": 0, "OTHER": 0}
    for section in checklist_data.get("checklist_sections", []) or []:
        if not isinstance(section, dict):
            continue
        for item in section.get("items", []) or []:
            if not isinstance(item, dict):
                continue
            status = str(item.get("status") or "VERIFY").strip().upper()
            if status not in counts:
                status = "OTHER"
            counts[status] += 1
    return counts


def build_site_survey_report_docx(checklist_data: dict[str, Any]) -> bytes | None:
    if not isinstance(checklist_data, dict):
        return None

    project_info = checklist_data.get("project_info") or {}
    project_name = str(project_info.get("project_name") or "AV Site Survey").strip()
    client = str(project_info.get("client") or "").strip()
    location = str(project_info.get("location") or "").strip()
    rooms = [
        str(item).strip()
        for item in checklist_data.get("rooms_areas", []) or []
        if str(item).strip()
    ]

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)

    cover = document.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cover.add_run("AV SITE SURVEY REPORT")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = __import__("docx").shared.RGBColor.from_string(NAVY)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(project_name)
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = __import__("docx").shared.RGBColor.from_string(BLUE)

    status_note = document.add_paragraph()
    status_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = status_note.add_run(
        "Draft / Pre-Survey Report – VERIFY items require confirmation during the site survey."
    )
    run.italic = True
    run.font.size = Pt(9)

    document.add_paragraph()

    info_table = document.add_table(rows=0, cols=2)
    info_table.style = "Table Grid"
    info_rows = [
        ("Project", project_name),
        ("Client", client),
        ("Location", location),
        ("Rooms / Areas", ", ".join(rooms)),
        ("Document Status", "Draft / Pre-Survey"),
    ]
    for label, value in info_rows:
        if not value:
            continue
        cells = info_table.add_row().cells
        _set_cell_text(cells[0], label, bold=True, color=WHITE)
        _shade_cell(cells[0], NAVY)
        _set_cell_text(cells[1], value)

    _add_section_heading(document, "1. Executive Summary")
    counts = _status_summary(checklist_data)
    summary_text = (
        f"This report has been prepared from the available project scope and survey checklist "
        f"for {project_name}. It is intended to organize site verification and close-out actions "
        f"before final handover or implementation. Current checklist status: "
        f"{counts['PASS']} PASS, {counts['VERIFY']} VERIFY, {counts['ACTION']} ACTION, "
        f"{counts['N/A']} N/A."
    )
    document.add_paragraph(summary_text)

    _add_section_heading(document, "2. Survey Scope")
    if rooms:
        document.add_paragraph(
            "The survey covers the following rooms / areas: " + ", ".join(rooms) + "."
        )
    else:
        document.add_paragraph(
            "Rooms / areas are to be confirmed against the approved project documentation."
        )
    document.add_paragraph(
        "All VERIFY items below require physical confirmation on site. The report must not be "
        "treated as confirmation of existing site conditions until those items are completed."
    )

    _add_section_heading(document, "3. Findings and Verification Status")
    findings = document.add_table(rows=1, cols=5)
    findings.style = "Table Grid"
    headers = ["Section", "Survey Item", "Status", "Notes / Evidence", "Action Required"]
    for index, header in enumerate(headers):
        _set_cell_text(findings.rows[0].cells[index], header, bold=True, color=WHITE)
        _shade_cell(findings.rows[0].cells[index], NAVY)

    for section_data in checklist_data.get("checklist_sections", []) or []:
        if not isinstance(section_data, dict):
            continue
        section_name = str(section_data.get("section") or "").strip()
        for item in section_data.get("items", []) or []:
            if not isinstance(item, dict):
                continue
            status = str(item.get("status") or "VERIFY").strip().upper()
            notes = str(item.get("notes") or "").strip()
            action = ""
            if status == "VERIFY":
                action = "Verify on site"
            elif status == "ACTION":
                action = "Corrective action / coordination required"
            row = findings.add_row().cells
            values = [
                section_name,
                item.get("item") or "",
                status,
                notes,
                action,
            ]
            for idx, value in enumerate(values):
                _set_cell_text(row[idx], value)
            if status == "VERIFY":
                _shade_cell(row[2], LIGHT_BLUE)
            elif status == "ACTION":
                _shade_cell(row[2], "FCE4D6")
            elif status == "PASS":
                _shade_cell(row[2], "E2F0D9")

    _add_section_heading(document, "4. Required Site Photos / Evidence")
    required_photos = checklist_data.get("required_photos", []) or []
    if required_photos:
        for item in required_photos:
            document.add_paragraph(f"☐ {item}")
    else:
        document.add_paragraph(
            "Photo evidence requirements are to be confirmed during the survey."
        )

    photo_register = checklist_data.get("photo_register", []) or []
    if photo_register:
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for idx, header in enumerate(["Photo", "Subject / Area", "Status", "Notes"]):
            _set_cell_text(table.rows[0].cells[idx], header, bold=True, color=WHITE)
            _shade_cell(table.rows[0].cells[idx], NAVY)

        for entry in photo_register:
            row = table.add_row().cells
            if isinstance(entry, dict):
                values = [
                    entry.get("photo_number") or "",
                    entry.get("photo_subject") or entry.get("subject") or "",
                    entry.get("status") or "PENDING",
                    entry.get("photo_notes") or entry.get("notes") or "",
                ]
            else:
                values = ["", str(entry), "PENDING", ""]
            for idx, value in enumerate(values):
                _set_cell_text(row[idx], value)

    _add_section_heading(document, "5. Open Items / Follow-up")
    open_items = checklist_data.get("open_items", []) or []
    if open_items:
        for item in open_items:
            document.add_paragraph(f"• {item}")
    else:
        document.add_paragraph(
            "No open items have been recorded yet. Update this section after site verification."
        )

    _add_section_heading(document, "6. Recommendations / Next Actions")
    document.add_paragraph(
        "• Complete all VERIFY items and attach supporting photo or measurement evidence where applicable."
    )
    document.add_paragraph(
        "• Convert unresolved items to ACTION with a named owner, target date, and closure evidence."
    )
    document.add_paragraph(
        "• Confirm any deviation from the design documentation before installation, programming, or handover."
    )
    document.add_paragraph(
        "• Issue the final report only after the observed site condition has been reconciled with the approved scope."
    )

    _add_section_heading(document, "7. Sign-off")
    sign_table = document.add_table(rows=4, cols=2)
    sign_table.style = "Table Grid"
    sign_rows = [
        ("Surveyed By", ""),
        ("Project / Engineering Review", ""),
        ("Client Representative", ""),
        ("Date", ""),
    ]
    for row, values in zip(sign_table.rows, sign_rows):
        _set_cell_text(row.cells[0], values[0], bold=True)
        _shade_cell(row.cells[0], LIGHT_GRAY)
        _set_cell_text(row.cells[1], values[1])

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()



def build_professional_site_survey_checklist_docx(
    checklist_data: dict[str, Any],
) -> bytes | None:
    if not isinstance(checklist_data, dict):
        return None

    project_info = checklist_data.get("project_info") or {}
    project_name = str(project_info.get("project_name") or "AV Site Survey").strip()
    client = str(project_info.get("client") or "").strip()
    location = str(project_info.get("location") or "").strip()
    rooms = [
        str(item).strip()
        for item in checklist_data.get("rooms_areas", []) or []
        if str(item).strip()
    ]

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(9.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("AV SITE SURVEY CHECKLIST")
    run.bold = True
    run.font.size = Pt(21)
    run.font.color.rgb = __import__("docx").shared.RGBColor.from_string(NAVY)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(project_name)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = __import__("docx").shared.RGBColor.from_string(BLUE)

    document.add_paragraph()

    info_table = document.add_table(rows=0, cols=2)
    info_table.style = "Table Grid"
    for label, value in [
        ("Project", project_name),
        ("Client", client),
        ("Location", location),
        ("Rooms / Areas", ", ".join(rooms)),
        ("Survey Status", "Pre-Survey / To Be Verified"),
    ]:
        if not value:
            continue
        cells = info_table.add_row().cells
        _set_cell_text(cells[0], label, bold=True, color=WHITE)
        _shade_cell(cells[0], NAVY)
        _set_cell_text(cells[1], value)

    document.add_paragraph(
        "Status guide: PASS = verified acceptable; VERIFY = confirm on site; "
        "ACTION = corrective action/coordination required; N/A = not applicable."
    )

    item_number = 1
    for section_data in checklist_data.get("checklist_sections", []) or []:
        if not isinstance(section_data, dict):
            continue

        section_name = str(section_data.get("section") or "").strip()
        items = section_data.get("items") or []
        if not section_name or not items:
            continue

        _add_section_heading(document, section_name)

        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for idx, header in enumerate(["No.", "Inspection / Verification Item", "Status", "Notes / Evidence", "Photo Ref."]):
            _set_cell_text(table.rows[0].cells[idx], header, bold=True, color=WHITE)
            _shade_cell(table.rows[0].cells[idx], NAVY)

        for item in items:
            if not isinstance(item, dict):
                continue
            status = str(item.get("status") or "VERIFY").strip().upper()
            row = table.add_row().cells
            values = [
                item_number,
                item.get("item") or "",
                status,
                item.get("notes") or "",
                item.get("photo_ref") or "",
            ]
            for idx, value in enumerate(values):
                _set_cell_text(row[idx], value)

            if status == "VERIFY":
                _shade_cell(row[2], LIGHT_BLUE)
            elif status == "ACTION":
                _shade_cell(row[2], "FCE4D6")
            elif status == "PASS":
                _shade_cell(row[2], "E2F0D9")

            item_number += 1

    _add_section_heading(document, "Required Site Photos")
    required_photos = checklist_data.get("required_photos", []) or []
    if required_photos:
        for photo in required_photos:
            document.add_paragraph(f"☐ {photo}")
    else:
        document.add_paragraph("☐ General room overview")
        document.add_paragraph("☐ Rack front and rear")
        document.add_paragraph("☐ Displays / camera / microphones / control interfaces")

    _add_section_heading(document, "Open Items / Follow-up")
    open_items = checklist_data.get("open_items", []) or []
    if open_items:
        for item in open_items:
            document.add_paragraph(f"☐ {item}")
    else:
        document.add_paragraph("☐ ________________________________________________")
        document.add_paragraph("☐ ________________________________________________")

    _add_section_heading(document, "Survey Sign-off")
    sign_table = document.add_table(rows=5, cols=2)
    sign_table.style = "Table Grid"
    sign_rows = [
        ("Surveyed By", ""),
        ("Project / Engineering Review", ""),
        ("Client Representative", ""),
        ("Date", ""),
        ("Final Status", "☐ Ready   ☐ Ready with Actions   ☐ Hold"),
    ]
    for row, values in zip(sign_table.rows, sign_rows):
        _set_cell_text(row.cells[0], values[0], bold=True)
        _shade_cell(row.cells[0], LIGHT_GRAY)
        _set_cell_text(row.cells[1], values[1])

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
