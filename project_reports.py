from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from project_dashboard import build_project_dashboard
from project_engineer import build_progress_report
from project_plan import build_phase_engineering_plan


REPORT_PERIODS = {"daily", "weekly", "monthly", "final"}


def _dedupe(values: list[str]) -> list[str]:
    output: list[str] = []
    seen = set()
    for raw in values:
        value = str(raw or "").strip()
        key = value.casefold()
        if value and key not in seen:
            output.append(value)
            seen.add(key)
    return output


def _flatten(entries: list[dict], key: str) -> list[str]:
    return _dedupe(
        [
            item
            for entry in entries or []
            for item in (entry.get(key) or [])
        ]
    )


def _project_identity(project: dict) -> dict:
    return {
        "project_name": project.get("name") or project.get("project_name"),
        "client": project.get("client"),
        "location": project.get("location"),
        "opportunity_number": project.get("opportunity_number"),
        "phase": project.get("phase"),
        "status": project.get("status"),
    }


def _latest_drawing(snapshot: dict) -> dict | None:
    drawings = snapshot.get("drawing_analyses") or []
    return drawings[0] if drawings else None


def _drawing_summary(latest: dict | None) -> dict:
    if not latest:
        return {
            "available": False,
            "file_name": None,
            "page_count": 0,
            "sheet_count": 0,
            "finding_count": 0,
            "high_findings": [],
            "connection_graph": {},
            "visual_reconciliation": {},
            "rooms": [],
        }

    qa = latest.get("qa") or {}
    findings = qa.get("findings") or []
    register = latest.get("register") or {}
    high = [
        finding
        for finding in findings
        if str(finding.get("severity") or "").lower() == "high"
    ]
    return {
        "available": True,
        "file_name": latest.get("file_name"),
        "page_count": latest.get("page_count") or 0,
        "sheet_count": register.get("sheet_count") or 0,
        "finding_count": len(findings),
        "high_findings": high[:12],
        "connection_graph": qa.get("connection_graph") or {},
        "visual_reconciliation": qa.get("visual_reconciliation") or {},
        "rooms": register.get("rooms") or [],
        "coordination_requirements": register.get("coordination_requirements") or [],
    }


def _executive_summary(
    *,
    identity: dict,
    period: str,
    dashboard: dict,
    progress: dict,
    drawing: dict,
) -> str:
    project_name = identity.get("project_name") or "the project"
    health = str(dashboard.get("health") or "unknown").replace("_", " ")
    completed = len(progress.get("completed") or [])
    issues = len(progress.get("issues") or [])
    blockers = len(progress.get("blockers") or [])
    high = len(drawing.get("high_findings") or [])
    graph = drawing.get("connection_graph") or {}
    resolved = int(graph.get("resolved_edge_count") or 0)
    ambiguous = int(graph.get("ambiguous_edge_count") or 0)

    if period == "final":
        opening = f"Final handover status review for {project_name}."
    else:
        opening = f"{period.capitalize()} project progress review for {project_name}."

    return (
        f"{opening} Current project health is {health}. "
        f"{completed} completed work item(s), {issues} issue(s), and {blockers} blocker(s) "
        f"are recorded for the reporting scope. "
        f"The latest drawing QA contains {high} high-priority engineering finding(s). "
        f"Connection review currently has {resolved} resolved and {ambiguous} ambiguous "
        f"connection(s). Items marked REVIEW or VERIFY require engineering or field confirmation "
        f"before dependent installation, programming, commissioning, or handover activities."
    )


def _final_handover_readiness(snapshot: dict, dashboard: dict, plan: dict) -> dict:
    blockers = dashboard.get("priority", {}).get("blockers") or []
    high_findings = dashboard.get("priority", {}).get("high_findings") or []
    open_issues = dashboard.get("priority", {}).get("issues") or []
    project = snapshot.get("project") or {}
    phase = project.get("phase")

    ready = (
        phase == "Handover"
        and not blockers
        and not high_findings
        and not open_issues
    )

    outstanding = []
    if phase != "Handover":
        outstanding.append(f"Project phase is currently '{phase}', not Handover.")
    if blockers:
        outstanding.append(f"{len(blockers)} blocker(s) remain open.")
    if high_findings:
        outstanding.append(
            f"{len(high_findings)} high-priority drawing QA finding(s) remain open."
        )
    if open_issues:
        outstanding.append(f"{len(open_issues)} issue(s) remain open.")

    outstanding.extend(plan.get("exit_criteria") or [])

    return {
        "status": "READY_FOR_FINAL_ACCEPTANCE" if ready else "NOT_READY_FOR_FINAL_ACCEPTANCE",
        "ready": ready,
        "outstanding": _dedupe(outstanding)[:20],
    }


def build_professional_project_report(
    snapshot: dict,
    *,
    period: str,
    anchor_date: date | None = None,
) -> dict:
    if period not in REPORT_PERIODS:
        raise ValueError("period must be daily, weekly, monthly, or final")

    project = snapshot.get("project") or {}
    identity = _project_identity(project)
    entries = snapshot.get("progress_entries") or []
    latest = _latest_drawing(snapshot)
    drawing = _drawing_summary(latest)
    dashboard = build_project_dashboard(snapshot)
    plan = build_phase_engineering_plan(snapshot)

    if period == "final":
        scoped_entries = entries
        progress = {
            "period": "final",
            "date_from": entries[0].get("date") if entries else None,
            "date_to": entries[-1].get("date") if entries else None,
            "entries_count": len(entries),
            "rooms": _flatten(entries, "rooms"),
            "completed": _flatten(entries, "completed"),
            "in_progress": _flatten(entries, "in_progress"),
            "issues": _flatten(entries, "issues"),
            "blockers": _flatten(entries, "blockers"),
            "next_actions": _flatten(entries, "next_actions"),
            "responsible_parties": _flatten(entries, "responsible_parties"),
            "entries": entries,
        }
    else:
        progress = build_progress_report(
            entries,
            period=period,
            anchor_date=anchor_date or date.today(),
            project=identity,
        )
        scoped_entries = progress.get("entries") or []
        progress["responsible_parties"] = _flatten(
            scoped_entries, "responsible_parties"
        )

    handover = _final_handover_readiness(snapshot, dashboard, plan)

    report = {
        "report_type": (
            "Final AV Project Handover Report"
            if period == "final"
            else f"{period.capitalize()} AV Project Progress Report"
        ),
        "period": period,
        "project": identity,
        "reporting_window": {
            "date_from": progress.get("date_from"),
            "date_to": progress.get("date_to"),
        },
        "executive_summary": "",
        "project_health": {
            "status": dashboard.get("health"),
            "phase": dashboard.get("phase"),
            "kpis": dashboard.get("kpis"),
        },
        "progress": {
            "entries_count": progress.get("entries_count") or 0,
            "rooms": progress.get("rooms") or [],
            "completed": progress.get("completed") or [],
            "in_progress": progress.get("in_progress") or [],
            "issues": progress.get("issues") or [],
            "blockers": progress.get("blockers") or [],
            "next_actions": progress.get("next_actions") or [],
            "responsible_parties": progress.get("responsible_parties") or [],
            "entries": scoped_entries,
        },
        "drawing_engineering_review": drawing,
        "engineering_plan": {
            "readiness": plan.get("readiness"),
            "objective": plan.get("objective"),
            "recommended_next_actions": plan.get("recommended_next_actions") or [],
            "exit_criteria": plan.get("exit_criteria") or [],
            "priority_findings": plan.get("priority_findings") or [],
        },
        "handover_readiness": handover,
        "disclaimer": (
            "This report is generated from the project records, uploaded drawings, "
            "document analysis, and logged progress available to the system. "
            "REVIEW/VERIFY findings require professional engineering, field, or "
            "manufacturer-specific confirmation before execution or acceptance."
        ),
    }

    report["executive_summary"] = _executive_summary(
        identity=identity,
        period=period,
        dashboard=dashboard,
        progress=progress,
        drawing=drawing,
    )
    return report


def _add_bullets(document: Document, items: list[str], empty_text: str) -> None:
    if not items:
        document.add_paragraph(empty_text)
        return
    for item in items:
        document.add_paragraph(str(item), style="List Bullet")


def _add_key_value_table(document: Document, values: list[tuple[str, Any]]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for key, value in values:
        cells = table.add_row().cells
        cells[0].text = str(key)
        cells[1].text = "" if value is None else str(value)


def build_professional_project_report_docx(report: dict) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(report.get("report_type") or "AV Project Report")
    run.bold = True
    run.font.size = Pt(18)

    project = report.get("project") or {}
    _add_key_value_table(
        document,
        [
            ("Project", project.get("project_name")),
            ("Client", project.get("client")),
            ("Location", project.get("location")),
            ("Opportunity / Project Ref.", project.get("opportunity_number")),
            ("Current Phase", project.get("phase")),
            ("Project Status", project.get("status")),
            (
                "Reporting Period",
                f"{(report.get('reporting_window') or {}).get('date_from') or '-'} "
                f"to {(report.get('reporting_window') or {}).get('date_to') or '-'}",
            ),
        ],
    )

    document.add_heading("1. Executive Summary", level=1)
    document.add_paragraph(report.get("executive_summary") or "")

    health = report.get("project_health") or {}
    kpis = health.get("kpis") or {}
    document.add_heading("2. Project Health & KPIs", level=1)
    _add_key_value_table(
        document,
        [
            ("Health", str(health.get("status") or "").replace("_", " ").upper()),
            ("Progress Updates", kpis.get("progress_updates", 0)),
            ("Open Issues", kpis.get("open_issues", 0)),
            ("Open Blockers", kpis.get("open_blockers", 0)),
            ("High QA Findings", kpis.get("high_findings", 0)),
            ("Graph Resolution", f"{kpis.get('graph_resolution_percent', 0)}%"),
            (
                "Visually Confirmed Connections",
                kpis.get("visual_confirmed_connections", 0),
            ),
            ("Visual Conflicts", kpis.get("visual_conflicts", 0)),
        ],
    )

    progress = report.get("progress") or {}
    document.add_heading("3. Progress Summary", level=1)
    document.add_heading("Completed Work", level=2)
    _add_bullets(document, progress.get("completed") or [], "No completed work recorded.")
    document.add_heading("Work in Progress", level=2)
    _add_bullets(document, progress.get("in_progress") or [], "No work in progress recorded.")
    document.add_heading("Issues / Snags", level=2)
    _add_bullets(document, progress.get("issues") or [], "No issues recorded.")
    document.add_heading("Blockers / Dependencies", level=2)
    _add_bullets(document, progress.get("blockers") or [], "No blockers recorded.")
    document.add_heading("Next Actions", level=2)
    _add_bullets(document, progress.get("next_actions") or [], "No next actions recorded.")

    drawing = report.get("drawing_engineering_review") or {}
    document.add_heading("4. Drawing & Engineering QA", level=1)
    if not drawing.get("available"):
        document.add_paragraph("No drawing analysis is currently available for this project.")
    else:
        graph = drawing.get("connection_graph") or {}
        visual = drawing.get("visual_reconciliation") or {}
        _add_key_value_table(
            document,
            [
                ("Latest Drawing", drawing.get("file_name")),
                ("PDF Pages", drawing.get("page_count", 0)),
                ("Indexed Sheets", drawing.get("sheet_count", 0)),
                ("QA Findings", drawing.get("finding_count", 0)),
                ("Resolved Connections", graph.get("resolved_edge_count", 0)),
                ("Ambiguous Connections", graph.get("ambiguous_edge_count", 0)),
                ("Visual Confirmations", visual.get("confirmed_count", 0)),
                ("Visual Conflicts", visual.get("conflict_count", 0)),
            ],
        )
        document.add_heading("High-Priority Findings", level=2)
        findings = drawing.get("high_findings") or []
        if findings:
            for finding in findings:
                paragraph = document.add_paragraph()
                paragraph.add_run(
                    f"{finding.get('title') or 'Engineering finding'}"
                ).bold = True
                action = finding.get("recommended_action")
                if action:
                    document.add_paragraph(f"Required action: {action}")
        else:
            document.add_paragraph("No high-priority drawing QA findings recorded.")

    plan = report.get("engineering_plan") or {}
    document.add_heading("5. Engineering Next Steps", level=1)
    document.add_paragraph(
        f"Readiness: {str(plan.get('readiness') or '').replace('_', ' ').upper()}"
    )
    if plan.get("objective"):
        document.add_paragraph(str(plan["objective"]))
    _add_bullets(
        document,
        plan.get("recommended_next_actions") or [],
        "No engineering next actions recorded.",
    )

    if report.get("period") == "final":
        handover = report.get("handover_readiness") or {}
        document.add_heading("6. Final Handover Readiness", level=1)
        document.add_paragraph(
            str(handover.get("status") or "").replace("_", " ")
        )
        _add_bullets(
            document,
            handover.get("outstanding") or [],
            "No outstanding acceptance items recorded.",
        )
        disclaimer_heading = "7. Report Basis & Limitations"
    else:
        disclaimer_heading = "6. Report Basis & Limitations"

    document.add_heading(disclaimer_heading, level=1)
    document.add_paragraph(report.get("disclaimer") or "")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = "AV Project Report · Generated from recorded project evidence"

    output = BytesIO()
    document.save(output)
    return output.getvalue()
