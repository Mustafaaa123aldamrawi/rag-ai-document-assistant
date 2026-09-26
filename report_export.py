from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _text(value: Any, fallback: str = "Not recorded") -> str:
    rendered = str(value or "").strip()
    return rendered or fallback


def _add_bullets(document: Document, items: list[str], empty: str = "No recorded items.") -> None:
    if not items:
        document.add_paragraph(empty)
        return
    for item in items:
        document.add_paragraph(str(item), style="List Bullet")


def _add_finding(document: Document, item: dict) -> None:
    title = _text(item.get("title"), "Engineering finding")
    severity = _text(item.get("severity"), "review").upper()
    status = _text(item.get("status"), "REVIEW").upper()

    p = document.add_paragraph()
    run = p.add_run(f"{severity} · {status} — {title}")
    run.bold = True

    why = str(item.get("why_it_matters") or "").strip()
    action = str(item.get("recommended_action") or "").strip()
    if why:
        document.add_paragraph(f"Why it matters: {why}")
    if action:
        document.add_paragraph(f"Recommended action: {action}")


def build_report_docx(payload: dict) -> bytes:
    document = Document()

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    title = _text(payload.get("title"), "AV Project Report")
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(20)

    project = payload.get("project") or {}
    period = payload.get("period") or {}

    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    project_rows = [
        ("Project", project.get("name")),
        ("Client", project.get("client")),
        ("Location", project.get("location")),
        ("Opportunity / Reference", project.get("opportunity_number")),
        ("Project Phase", project.get("phase")),
        ("Project Status", project.get("status")),
        ("Report From", period.get("date_from")),
        ("Report To", period.get("date_to")),
        ("Report Anchor Date", period.get("anchor_date")),
    ]
    for label, value in project_rows:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = _text(value)

    document.add_heading("Executive Summary", level=1)
    executive = payload.get("executive_summary") or {}
    document.add_paragraph(_text(executive.get("summary"), "No executive summary recorded."))
    if executive.get("project_health"):
        document.add_paragraph(
            f"Project health: {_text(executive.get('project_health')).replace('_', ' ').upper()}"
        )

    progress = payload.get("progress") or {}
    progress_sections = (
        ("Rooms / Areas", progress.get("rooms_areas") or []),
        ("Completed Work", progress.get("completed") or []),
        ("Work in Progress", progress.get("in_progress") or []),
        ("Issues / Snags", progress.get("issues_snags") or []),
        ("Blockers / Dependencies", progress.get("blockers_dependencies") or []),
        ("Next Actions", progress.get("next_actions") or []),
        ("Responsible Parties", progress.get("responsible_parties") or []),
    )
    for heading_text, items in progress_sections:
        document.add_heading(heading_text, level=1)
        _add_bullets(document, items)

    engineering = payload.get("engineering_review") or {}
    document.add_heading("Drawing & Engineering QA", level=1)

    latest = engineering.get("latest_drawing") or {}
    if latest:
        document.add_paragraph(
            f"Latest drawing set: {_text(latest.get('file_name'))} "
            f"({_text(latest.get('page_count'), '0')} pages)"
        )

    graph = engineering.get("connection_graph") or {}
    graph_table = document.add_table(rows=1, cols=5)
    graph_table.style = "Table Grid"
    graph_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Devices", "Wires", "Resolved", "Verify", "Resolution"]
    for idx, value in enumerate(headers):
        graph_table.rows[0].cells[idx].text = value
    values = [
        graph.get("devices", 0),
        graph.get("wires", 0),
        graph.get("resolved", 0),
        graph.get("ambiguous", 0),
        f"{graph.get('resolution_percent', 0)}%",
    ]
    row = graph_table.add_row().cells
    for idx, value in enumerate(values):
        row[idx].text = str(value)

    visual = engineering.get("visual_reconciliation") or {}
    document.add_paragraph(
        "Visual reconciliation: "
        f"{visual.get('confirmed', 0)} confirmed, "
        f"{visual.get('conflicts', 0)} conflict(s), "
        f"{visual.get('no_visual_match', 0)} without visual match, "
        f"{visual.get('visual_only', 0)} visual-only connection(s)."
    )

    high_findings = engineering.get("high_findings") or []
    document.add_heading("Priority Engineering Findings", level=2)
    if high_findings:
        for item in high_findings:
            _add_finding(document, item)
    else:
        document.add_paragraph("No high-severity engineering findings recorded.")

    management = payload.get("management_summary") or {}
    document.add_heading("Management Summary", level=1)
    _add_bullets(
        document,
        management.get("priority_blockers") or [],
        "No priority blockers recorded.",
    )
    if management.get("priority_issues"):
        p = document.add_paragraph()
        p.add_run("Priority issues").bold = True
        _add_bullets(document, management.get("priority_issues") or [])
    if management.get("priority_next_actions"):
        p = document.add_paragraph()
        p.add_run("Priority next actions").bold = True
        _add_bullets(document, management.get("priority_next_actions") or [])

    readiness = payload.get("handover_readiness")
    if readiness:
        document.add_heading("Handover Readiness", level=1)
        document.add_paragraph(
            f"Status: {_text(readiness.get('status')).replace('_', ' ').upper()}"
        )
        gate_table = document.add_table(rows=1, cols=2)
        gate_table.style = "Table Grid"
        gate_table.rows[0].cells[0].text = "Check"
        gate_table.rows[0].cells[1].text = "Status"
        for check in readiness.get("checks") or []:
            cells = gate_table.add_row().cells
            cells[0].text = _text(check.get("label"))
            cells[1].text = "PASS" if check.get("passed") else "OPEN"
        note = str(readiness.get("note") or "").strip()
        if note:
            document.add_paragraph(note)

    document.add_heading("Quality Note", level=1)
    document.add_paragraph(
        _text(
            payload.get("quality_statement"),
            (
                "Engineering review items must be validated against approved drawings, "
                "manufacturer documentation, site conditions, and commissioning evidence."
            ),
        )
    )

    document.add_paragraph()
    footer = document.add_paragraph(
        "Generated by AV Intelligence Assistant"
    )
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def safe_report_filename(payload: dict, extension: str = "docx") -> str:
    project = (payload.get("project") or {}).get("name") or "AV_Project"
    title = payload.get("report_type") or "report"
    value = f"{project}_{title}".strip()
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", value)
    value = re.sub(r"_+", "_", value).strip("._-")
    return f"{value or 'AV_Project_Report'}.{extension.lstrip('.')}"
