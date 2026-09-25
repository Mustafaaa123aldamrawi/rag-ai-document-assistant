from io import BytesIO

from docx import Document

from deliverables import build_professional_site_survey_checklist_docx, build_site_survey_report_docx


def sample_data():
    return {
        "project_info": {
            "project_name": "Mastercard Riyadh",
            "client": "Mastercard",
            "location": "Riyadh",
        },
        "rooms_areas": ["Room-1", "Room-2"],
        "checklist_sections": [
            {
                "section": "AV Rack",
                "items": [
                    {
                        "item": "Verify rack capacity",
                        "status": "VERIFY",
                        "notes": "",
                    },
                    {
                        "item": "Existing PDU retained",
                        "status": "PASS",
                        "notes": "Visible in scope.",
                    },
                ],
            }
        ],
        "required_photos": ["Rack front", "Rack rear"],
        "open_items": ["Confirm VLAN requirements"],
    }


def test_report_docx_is_valid_word_file():
    output = build_site_survey_report_docx(sample_data())
    assert output
    assert output[:2] == b"PK"


def test_report_contains_professional_sections():
    output = build_site_survey_report_docx(sample_data())
    document = Document(BytesIO(output))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "AV SITE SURVEY REPORT" in text
    assert "Executive Summary" in text
    assert "Recommendations / Next Actions" in text
    assert "Draft / Pre-Survey Report" in text



def test_professional_checklist_docx_is_valid():
    output = build_professional_site_survey_checklist_docx(sample_data())
    assert output
    document = Document(BytesIO(output))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "AV SITE SURVEY CHECKLIST" in text
    assert "Survey Sign-off" in text



def visual_sample_data():
    data = sample_data()
    data["executive_summary"] = (
        "Two site photos were reviewed. One possible visual issue requires field verification."
    )
    data["inspection_meta"] = {
        "overall_status": "ACTION REQUIRED",
        "photos_reviewed": 2,
        "possible_issues": 1,
        "verify_items": 1,
    }
    data["visual_inspection"] = {
        "inspection_meta": data["inspection_meta"],
        "executive_summary": data["executive_summary"],
        "visual_findings": [
            {
                "id": "F-01",
                "source_photo": "rack.jpg",
                "finding": "A visible cable appears partially disconnected.",
                "basis": "The connector is visibly not fully seated.",
                "confidence": "MEDIUM",
                "priority": "MEDIUM",
                "status": "OPEN",
                "required_action": "Verify the connection on site before assigning root cause.",
            }
        ],
        "verification_items": [
            {
                "id": "V-01",
                "source_photo": "display.jpg",
                "item": "Display manufacturer is not readable.",
                "status": "VERIFY",
                "required_action": "Verify from the device label or management interface.",
            }
        ],
    }
    data["visual_findings"] = data["visual_inspection"]["visual_findings"]
    data["verification_items"] = data["visual_inspection"]["verification_items"]
    data["photo_register"] = [
        {
            "photo_number": "Photo 01",
            "file_name": "display.jpg",
            "category": "SITE_PHOTO",
            "subject_equipment": "Curved display",
            "status": "REVIEWED",
            "notes": "Display is active.",
        },
        {
            "photo_number": "Photo 02",
            "file_name": "rack.jpg",
            "category": "AV_EQUIPMENT",
            "subject_equipment": "AV rack",
            "status": "REVIEWED",
            "notes": "Rack equipment is visible.",
        },
    ]
    return data


def test_report_includes_visual_findings_and_photo_register():
    output = build_site_survey_report_docx(visual_sample_data())
    document = Document(BytesIO(output))
    paragraph_text = "\n".join(p.text for p in document.paragraphs)
    table_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    combined = paragraph_text + "\n" + table_text

    assert "Visual Findings & Field Verification" in combined
    assert "Photo Evidence Register" in combined
    assert "ACTION REQUIRED" in combined
    assert "rack.jpg" in combined
    assert "display.jpg" in combined


def test_checklist_includes_visual_evidence_section():
    output = build_professional_site_survey_checklist_docx(visual_sample_data())
    document = Document(BytesIO(output))
    combined = "\n".join(
        [p.text for p in document.paragraphs]
        + [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
    )

    assert "Visual Findings & Field Verification" in combined
    assert "Photo Evidence Register" in combined
    assert "No." in combined



def test_client_facing_documents_use_compact_photo_refs_and_no_app_branding():
    data = visual_sample_data()
    data["photo_register"][0]["photo_ref"] = "P01"
    data["photo_register"][1]["photo_ref"] = "P02"
    data["visual_findings"][0]["source_photo_ref"] = "P02"
    data["visual_findings"][0]["status"] = "OBSERVATION"
    data["visual_inspection"]["visual_findings"] = data["visual_findings"]

    report = Document(BytesIO(build_site_survey_report_docx(data)))
    checklist = Document(BytesIO(build_professional_site_survey_checklist_docx(data)))

    report_text = "\n".join(
        [p.text for p in report.paragraphs]
        + [p.text for section in report.sections for p in section.footer.paragraphs]
        + [
            cell.text
            for table in report.tables
            for row in table.rows
            for cell in row.cells
        ]
    )
    checklist_text = "\n".join(
        [p.text for p in checklist.paragraphs]
        + [p.text for section in checklist.sections for p in section.footer.paragraphs]
        + [
            cell.text
            for table in checklist.tables
            for row in table.rows
            for cell in row.cells
        ]
    )

    assert "AV Intelligence Assistant" not in report_text
    assert "AV Intelligence Assistant" not in checklist_text
    assert "P02" in report_text
    assert "OBSERVATION" in report_text
    assert "AV SITE SURVEY REPORT & INSPECTION" in report_text


def test_visual_findings_are_not_repeated_inside_scope_table():
    data = visual_sample_data()
    finding_text = data["visual_findings"][0]["finding"]
    output = build_site_survey_report_docx(data)
    document = Document(BytesIO(output))
    combined = "\n".join(
        [p.text for p in document.paragraphs]
        + [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
    )
    assert combined.count(finding_text) == 1



def test_word_outputs_separate_survey_and_visual_status():
    data = visual_sample_data()
    data["survey_status"] = "FIELD VERIFICATION REQUIRED"
    data["visual_status"] = "NO OBVIOUS VISUAL FAULT"
    data["visual_inspection"]["inspection_meta"]["visual_status"] = "NO OBVIOUS VISUAL FAULT"

    report = Document(BytesIO(build_site_survey_report_docx(data)))
    checklist = Document(BytesIO(build_professional_site_survey_checklist_docx(data)))

    report_text = "\n".join(
        [p.text for p in report.paragraphs]
        + [
            cell.text
            for table in report.tables
            for row in table.rows
            for cell in row.cells
        ]
    )
    checklist_text = "\n".join(
        [p.text for p in checklist.paragraphs]
        + [
            cell.text
            for table in checklist.tables
            for row in table.rows
            for cell in row.cells
        ]
    )

    assert "FIELD VERIFICATION REQUIRED" in report_text
    assert "NO OBVIOUS VISUAL FAULT" in report_text
    assert "FIELD VERIFICATION REQUIRED" in checklist_text
    assert "Visual status: NO OBVIOUS VISUAL FAULT" in checklist_text
