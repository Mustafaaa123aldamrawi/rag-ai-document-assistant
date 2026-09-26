from datetime import date
from io import BytesIO

from docx import Document

from project_reports import (
    build_professional_project_report,
    build_professional_project_report_docx,
)


def sample_snapshot():
    return {
        "project": {
            "id": "p1",
            "name": "Pinsent Masons - Riyadh",
            "client": "Pinsent Masons",
            "location": "Riyadh, KSA",
            "opportunity_number": "390I-25-88927",
            "phase": "First Fix",
            "status": "active",
        },
        "progress_entries": [
            {
                "id": "e1",
                "date": "2026-09-26",
                "phase": "First Fix",
                "rooms": ["L7-07 Client Meeting 03"],
                "completed": ["Containment checked"],
                "in_progress": ["Cable pulling"],
                "issues": ["Data outlet not ready"],
                "blockers": ["Awaiting IT"],
                "next_actions": ["Pull cables after IT release"],
                "responsible_parties": ["IT"],
            }
        ],
        "drawing_analyses": [
            {
                "id": "d1",
                "file_name": "shop-drawings.pdf",
                "page_count": 25,
                "register": {
                    "sheet_count": 25,
                    "rooms": ["L7-07 Client Meeting 03"],
                    "coordination_requirements": [],
                },
                "qa": {
                    "findings": [
                        {
                            "severity": "high",
                            "status": "REVIEW",
                            "title": "V0001 endpoint conflict",
                            "recommended_action": "Verify both endpoints visually.",
                        }
                    ],
                    "connection_graph": {
                        "node_count": 6,
                        "edge_count": 5,
                        "resolved_edge_count": 4,
                        "ambiguous_edge_count": 1,
                        "resolution_rate": 0.8,
                    },
                    "visual_reconciliation": {
                        "confirmed_count": 3,
                        "conflict_count": 1,
                    },
                },
            }
        ],
        "reports": [],
    }


def test_professional_weekly_report_includes_engineering_context():
    report = build_professional_project_report(
        sample_snapshot(),
        period="weekly",
        anchor_date=date(2026, 9, 26),
    )

    assert report["report_type"] == "Weekly AV Project Progress Report"
    assert report["project"]["project_name"] == "Pinsent Masons - Riyadh"
    assert report["progress"]["completed"] == ["Containment checked"]
    assert report["progress"]["blockers"] == ["Awaiting IT"]
    assert report["drawing_engineering_review"]["finding_count"] == 1
    assert report["drawing_engineering_review"]["connection_graph"]["resolved_edge_count"] == 4
    assert "engineering review required" in report["executive_summary"].lower()


def test_final_report_not_ready_when_open_project_risks_exist():
    report = build_professional_project_report(
        sample_snapshot(),
        period="final",
    )

    assert report["report_type"] == "Final AV Project Handover Report"
    assert report["handover_readiness"]["ready"] is False
    assert (
        report["handover_readiness"]["status"]
        == "NOT_READY_FOR_FINAL_ACCEPTANCE"
    )
    assert report["handover_readiness"]["outstanding"]


def test_professional_report_docx_contains_client_ready_sections():
    report = build_professional_project_report(
        sample_snapshot(),
        period="weekly",
        anchor_date=date(2026, 9, 26),
    )
    payload = build_professional_project_report_docx(report)

    assert payload[:2] == b"PK"
    document = Document(BytesIO(payload))
    combined = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
    )

    assert "Weekly AV Project Progress Report" in combined
    assert "Executive Summary" in combined
    assert "Project Health & KPIs" in combined
    assert "Drawing & Engineering QA" in combined
    assert "Engineering Next Steps" in combined
    assert "Awaiting IT" in combined
    assert "V0001 endpoint conflict" in combined


def test_final_docx_contains_handover_readiness_section():
    report = build_professional_project_report(
        sample_snapshot(),
        period="final",
    )
    payload = build_professional_project_report_docx(report)
    document = Document(BytesIO(payload))
    combined = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
    )

    assert "Final AV Project Handover Report" in combined
    assert "Final Handover Readiness" in combined
    assert "NOT READY FOR FINAL ACCEPTANCE" in combined


def test_clean_handover_project_can_be_ready_for_final_acceptance():
    snapshot = sample_snapshot()
    snapshot["project"]["phase"] = "Handover"
    snapshot["progress_entries"] = [
        {
            "id": "e2",
            "date": "2026-09-26",
            "phase": "Handover",
            "completed": ["Client training complete"],
            "issues": [],
            "blockers": [],
            "next_actions": [],
        }
    ]
    snapshot["drawing_analyses"][0]["qa"]["findings"] = []
    snapshot["drawing_analyses"][0]["qa"]["visual_reconciliation"] = {
        "confirmed_count": 5,
        "conflict_count": 0,
    }

    report = build_professional_project_report(snapshot, period="final")

    assert report["handover_readiness"]["ready"] is True
    assert (
        report["handover_readiness"]["status"]
        == "READY_FOR_FINAL_ACCEPTANCE"
    )
